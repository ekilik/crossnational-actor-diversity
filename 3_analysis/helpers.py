"""Shared constants and helpers for the actor-diversity analysis pipeline.

Single source of helper functions for the actor-diversity notebooks (01_clean_corpus -> 05).
Country/topic/outlet labels, file paths, dedup parameters, and the small utility
functions live here so the notebooks stay DRY: change a label or a threshold once
and every notebook that imports it picks the change up. Avoids the copy-paste
drift that crept into the original scripts (BRAND_RULES, COUNTRY_ORDER, the
diversity formulas were each duplicated across notebooks).

Import convention in the notebooks:  import helpers as h
"""
from pathlib import Path
import re

import numpy as np
import pandas as pd
import warnings

# --- project paths -----------------------------------------------------------
# Resolved from THIS file, not the working directory, so a notebook works the
# same whether it is run from scripts/ or the project root.
BASE_DIR    = Path(__file__).resolve().parent.parent
DATA_DIR    = BASE_DIR / "analysis_data"
RESULTS_DIR = BASE_DIR / "results"
MODEL_DIR   = BASE_DIR / "model_dfs"

SOCIAL_RAW    = DATA_DIR / "df_actors_socialdiv_analysis.csv"
POLITICAL_RAW = DATA_DIR / "df_actors_politicaldiv_analysis.csv"
SOCIAL_CLEAN    = DATA_DIR / "df_actors_socialdiv_analysis_clean.csv"
POLITICAL_CLEAN = DATA_DIR / "df_actors_politicaldiv_analysis_clean.csv"

CSV_SEP = ";"   # all project CSVs are semicolon-separated

# --- country -----------------------------------------------------------------
COUNTRY_ORDER = ["BE", "IR", "NL", "TR", "UK"]
COUNTRY_NAME  = {"BE": "Belgium", "IR": "Ireland", "NL": "Netherlands",
                 "TR": "Turkey", "UK": "United Kingdom"}
# The first digit of an article_id encodes the country. Used to recover the
# country for political "orphan" rows whose articles are absent from the
# social df (and therefore carry no country of their own).
COUNTRY_BY_ID_PREFIX = {"1": "UK", "2": "IR", "3": "NL", "4": "BE", "5": "TR"}

# --- topic -------------------------------------------------------------------
# Canonical topic codes in the data are 'immigration' and 'rising_prices'.
# TOPIC_LABEL maps them to the display label used in tables/figures.
TOPIC_LABEL = {"rising_prices": "rising prices", "immigration": "immigration"}

# --- outlet type / display ---------------------------------------------------
TYPE_ORDER = ["Quality newspaper", "Popular newspaper", "Public broadcaster"]
TYPE_LABEL = {"Quality newspaper": "Upmarket", "Popular newspaper": "Mass-market",
              "Public broadcaster": "Public service"}
BRAND_DISP = {"RTE": "RTÉ", "Telegraaf": "De Telegraaf", "Volkskrant": "De Volkskrant",
              "The Daily Telegraph": "Daily Telegraph"}
LEANING = {
    "De Morgen": "Left",        "De Standaard": "Right",        # BE
    "The Irish Times": "Left",  "Irish Independent": "Right",   # IR
    "Volkskrant": "Left",       "NRC": "Right",                 # NL
    "Cumhuriyet": "Left",       "Sabah": "Right",               # TR
    "The Guardian": "Left",     "The Daily Telegraph": "Right", # UK
}

# --- political party ideology (left/right) -----------------------------------
# lrgen_cat is the 3-way expert placement (left/center/right). The cross-country
# IDEOLOGICAL diversity score needs a binary left/right variable (lrgen_cat2), so
# the parties coded 'center' are each assigned to a side (or NaN when genuinely
# unplaceable). TR and UK have no 'center' parties, so lrgen_cat2 == lrgen_cat
# there. Keyed by CANONICAL party_final, so apply PARTY_ALIASES first; parties
# absent from LRGEN2_RECODE keep their lrgen_cat value.
PARTY_ALIASES = {           # spelling variants of the SAME party -> canonical label
    "Les Engagés": "LE",    # BE: cdH renamed to Les Engagés in 2022
    "50PLUS": "50+",        # NL: pensioners' party, recorded under two stylings
}
LRGEN2_RECODE = {           # (country, party_final) -> 'left' | 'right' | np.nan
    ("BE", "CD&V"): "right", ("BE", "LE"):  "right",
    ("BE", "DéFI"): "right", ("BE", "Vooruit"): "left",
    ("NL", "CDA"):  "right", ("NL", "CU"):  "right",
    ("NL", "D66"):  "left",  ("NL", "Volt"): "left",
    ("NL", "50+"):  np.nan,  # single-issue pensioners' party, not placed on L/R
    ("IR", "FF"):   "right",
}
LRGEN2_ORDER = ["left", "right"]

# --- political-independence risk tier (MPM, for the descriptive H1/H2 grouping) ---
# Low = BE/NL/IR (19-26%), Medium = UK (41%), High = TR (89%). Note Medium and High
# are single countries, so cross-tier comparisons are descriptive, not inferential.
RISK_GROUP = {"BE": "Low", "NL": "Low", "IR": "Low", "UK": "Medium", "TR": "High"}
RISK_ORDER = ["Low", "Medium", "High"]

# --- societal actor categories (fixed set, S = 4) ----------------------------
CAT_COLS = ["nr_politicians", "nr_experts_professionals",
            "nr_civil_society", "nr_general_public"]

# --- deduplication parameters ------------------------------------------------
SCAN_THRESHOLD     = 0.90   # candidate near-duplicate pairs collected at/above this (calibration)
TEXT_SIM_THRESHOLD = 0.98   # DECISION cutoff: text pairs with cosine >= this are duplicates
SAME_BRAND_ONLY    = True   # only collapse near-duplicates that share an outlet


def country_from_id(article_id) -> str | None:
    """Recover the country code from the leading digit of an article_id."""
    return COUNTRY_BY_ID_PREFIX.get(str(article_id)[:1])


def parse_dates_safe(s: pd.Series, *, assert_complete: bool = True) -> pd.Series:
    """Parse a date Series robustly, failing loudly instead of dropping rows.

    The topic files mix ISO dates ('2024-11-29', from NexisUni/Belgapress) with
    human-readable scraped dates ('Friday, 29 Nov 2024 09:58', from broadcaster
    sites such as RTÉ). A fixed format or dayfirst=True with errors='coerce'
    silently turns the un-matched style into NaT; that is how 331 Irish RTÉ
    campaign-period articles were dropped previously. format='mixed' handles
    both. We then assert that nothing parsed to NaT that was not already missing,
    so a new/unexpected format raises rather than vanishing.
    """
    s = pd.Series(s)
    out = pd.to_datetime(s, format="mixed", errors="coerce")
    if assert_complete:
        newly_bad = out.isna() & s.notna() & (s.astype(str).str.strip() != "")
        n_bad = int(newly_bad.sum())
        if n_bad:
            examples = list(s[newly_bad].astype(str).unique()[:5])
            raise ValueError(
                f"{n_bad} dates failed to parse (and were not originally missing). "
                f"Examples: {examples}. Inspect the source rather than coercing them away."
            )
    return out


def load_topic_files(with_text: bool = True) -> pd.DataFrame:
    """Load the per-country topic files into one frame.

    Columns: article_id, country, topic, relevant, word_count, title[, text].
    These files carry the article body text (absent from the actor dfs, but
    needed for near-duplicate detection) and the relevance flag used for the
    corpus funnel. Country comes from the filename prefix, topic from the
    'immigration'/'inflation' marker in the filename.
    """
    frames = []
    for path in sorted(DATA_DIR.glob("*_df*.csv")):
        if "actors" in path.name:
            continue
        d = pd.read_csv(path, sep=CSV_SEP)
        cc = path.name[:2].upper()
        topic = "immigration" if "immigration" in path.name else "rising_prices"
        rel_col = next((c for c in d.columns if c.startswith("relevant_")), None)
        out = pd.DataFrame({
            "article_id": d["article_id"],
            "country": cc,
            "topic": topic,
            "relevant": pd.to_numeric(d[rel_col], errors="coerce") if rel_col else np.nan,
            "word_count": d.get("word_count"),
            "title": d["title"].astype(str) if "title" in d.columns else "",
        })
        if with_text:
            out["text"] = d["text"].astype(str) if "text" in d.columns else ""
        frames.append(out)
    return pd.concat(frames, ignore_index=True)


def norm_text(s) -> str:
    """Lowercase and collapse whitespace, for text-similarity comparison."""
    return re.sub(r"\s+", " ", str(s).strip().lower())


def week_start(year_week) -> pd.Series:
    """Monday date of an ISO 'YYYY-WW' week label (e.g. '2024-21' -> 2024-05-20)."""
    s = pd.Series(year_week).astype(str)
    return pd.to_datetime(s + "-1", format="%G-%V-%u")


# ---------------------------------------------------------------------------
# APA7 output helpers (figures + Word tables)
# ---------------------------------------------------------------------------
APA_RC = {
    "figure.dpi": 300, "savefig.dpi": 300,
    "font.family": "sans-serif", "font.sans-serif": ["Arial", "DejaVu Sans"],
    "font.size": 11, "axes.titlesize": 12, "axes.labelsize": 11,
    "axes.titleweight": "normal", "axes.edgecolor": "black", "axes.linewidth": 0.8,
    "axes.spines.top": False, "axes.spines.right": False, "axes.grid": False,
    "legend.frameon": False, "xtick.color": "black", "ytick.color": "black",
    "axes.labelcolor": "black", "text.color": "black",
}

def set_apa_style():
    """Apply APA7-friendly matplotlib defaults (sans-serif, no chartjunk, 300 dpi)."""
    import matplotlib.pyplot as plt
    plt.rcParams.update(APA_RC)

FIG_WIDTH_IN      = 6.5     # APA: US Letter minus 1in margins. A4 with 2.54cm margins = 6.27.
FIG_MAX_HEIGHT_IN = 7.5     # 9in text height, minus room for the number, title and note
_FIG_PAD_IN       = 0.02    # whitespace kept inside the saved bbox

def save_apa(fig, path, width_in=FIG_WIDTH_IN, dpi=300,
             max_height_in=FIG_MAX_HEIGHT_IN):
    """Save a figure at an exact physical width, so Word inserts it 1:1.

    bbox_inches="tight" crops to the drawn content, which makes the saved size an
    *output* of the layout rather than something you declare: the same figsize
    yields different widths depending on how long the legend labels are. Here the
    width is an input instead. Two details the naive version gets wrong: savefig
    adds pad_inches on each side after cropping (so the target must subtract it),
    and set_size_inches leaves font sizes fixed in points, so the tight bbox is
    not linear in the figure width. It is monotone, so a few passes converge.
    """
    target = width_in - 2 * _FIG_PAD_IN
    for _ in range(5):
        fig.canvas.draw()
        bb = fig.get_tightbbox(fig.canvas.get_renderer())
        if abs(bb.width - target) < 0.005:
            break
        w, hgt = fig.get_size_inches()
        fig.set_size_inches(w * target / bb.width, hgt)
        fig.savefig(path, dpi=dpi, bbox_inches="tight", pad_inches=_FIG_PAD_IN)
    saved_h = fig.get_tightbbox(fig.canvas.get_renderer()).height + 2 * _FIG_PAD_IN
    if saved_h > max_height_in:
        warnings.warn(
            f"{Path(path).name}: {saved_h:.2f}in tall, over the {max_height_in}in "
            f"page budget. Reduce the figsize height; letting Word scale it down "
            f"would take the labels below APA's 8pt minimum.",
            stacklevel=2)

    


def apa_p(p) -> str:
    """Format a p value APA-style: '< .001', else no leading zero ('.013')."""
    if p != p:           # NaN
        return ""
    return "< .001" if p < .001 else f"{p:.3f}".lstrip("0")


def apa_table(df: pd.DataFrame, number, title: str, name: str, note: str = "",
              index_label=None, spanners=None, landscape=False) -> str:
    """Write a table in both output formats: an APA7 Word table plus a CSV companion.

    `name` is the file stem: one call writes results/tables/docx/<name>.docx and
    results/tables/csv/<name>.csv from the same DataFrame. A single writer keeps the
    pair in sync by construction; no call site can produce one format and forget the
    other, or let the two file names drift apart.

    Word side: Times New Roman, horizontal rules only (top, below the column headers,
    bottom), plain table number line, italic title, italic 'Note.' prefix. Values are
    written as-is, so pre-format numbers before calling. `spanners` adds a top header
    row of column-group labels, e.g. [("", 2), ("Societal actors", 6)] (label, number
    of columns); a rule is drawn beneath each labelled span. The CSV carries the same
    pre-formatted values, including the index when `index_label` is set.
    """
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    docx_dir = RESULTS_DIR / "tables/docx"
    csv_dir  = RESULTS_DIR / "tables/csv"
    docx_dir.mkdir(parents=True, exist_ok=True)
    csv_dir.mkdir(parents=True, exist_ok=True)
    if index_label is not None:
        df.to_csv(csv_dir / f"{name}.csv", index_label=index_label)
    else:
        df.to_csv(csv_dir / f"{name}.csv", index=False)

    def rule(cell, edge, sz=6):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = tcPr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
        e = borders.find(qn(f"w:{edge}"))
        if e is None:
            e = OxmlElement(f"w:{edge}"); borders.append(e)
        for k, v in (("val", "single"), ("sz", str(sz)), ("color", "000000")):
            e.set(qn(f"w:{k}"), v)

    doc = Document()
    if landscape:
        from docx.enum.section import WD_ORIENT
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph(f"Table {number}")
    doc.add_paragraph().add_run(title).italic = True

    cols = ([index_label] if index_label is not None else []) + [str(c) for c in df.columns]
    n_head = 2 if spanners else 1
    hr = n_head - 1                       # row index of the column-name header
    tbl = doc.add_table(rows=len(df) + n_head, cols=len(cols))

    if spanners:
        j = 0
        for label, span in spanners:
            cell = tbl.cell(0, j).merge(tbl.cell(0, j + span - 1)) if span > 1 else tbl.cell(0, j)
            if label:
                cell.paragraphs[0].add_run(label)
                for jj in range(j, j + span):
                    rule(tbl.cell(0, jj), "bottom")
            j += span

    for j, c in enumerate(cols):
        tbl.cell(hr, j).paragraphs[0].add_run(str(c))
    for i, (idx, row) in enumerate(df.iterrows(), start=n_head):
        vals = ([str(idx)] if index_label is not None else []) + [str(v) for v in row]
        for j, v in enumerate(vals):
            tbl.cell(i, j).paragraphs[0].add_run(v)

    last = len(df) + n_head - 1
    for j in range(len(cols)):
        rule(tbl.cell(0, j), "top"); rule(tbl.cell(hr, j), "bottom"); rule(tbl.cell(last, j), "bottom")
    if note:
        p = doc.add_paragraph(); p.add_run("Note. ").italic = True; p.add_run(note)
    doc.save(str(docx_dir / f"{name}.docx"))
    return str(docx_dir / f"{name}.docx")


# ---------------------------------------------------------------------------
# Diversity indices (pooled operationalization)
# ---------------------------------------------------------------------------
def diversity_from_counts(counts) -> tuple:
    """Standardized Simpson's Dz for pooled category-count rows.

    counts: array (n_grains, S) of non-negative pooled counts; each row is one
    distribution (e.g. an outlet-week's four societal categories). Returns
    (Dz, n_total); grains with zero total are NaN.

    Dz = (1 - sum p_i^2) / (1 - 1/S)  -- standardized Gini-Simpson, 0..1
    (0 = one category dominates, 1 = all S categories equal). Counts are summed
    within the grain first, then indexed once, so zero-/single-actor articles add
    no/one actor to the pool rather than yielding undefined per-article values.
    """
    counts = np.asarray(counts, dtype=float)
    n = counts.sum(axis=1)
    S = counts.shape[1]
    with np.errstate(divide="ignore", invalid="ignore"):
        p = np.where(n[:, None] > 0, counts / n[:, None], 0.0)
        lam = (p ** 2).sum(axis=1)
    Dz = (1.0 - lam) / (1.0 - 1.0 / S)
    Dz[n == 0] = np.nan
    return Dz, n


def societal_dz(soc: pd.DataFrame, group_keys: list) -> pd.DataFrame:
    """Societal Simpson's Dz (four actor categories) per grain.

    Sums the category counts over group_keys, then computes Dz on the pooled
    distribution. Returns group_keys + n_actors + simpson_Dz.
    """
    counts = soc.groupby(group_keys, observed=True)[CAT_COLS].sum().reset_index()
    Dz, n = diversity_from_counts(counts[CAT_COLS].to_numpy())
    counts = counts.drop(columns=CAT_COLS)
    counts["n_actors"] = n.astype(int)
    counts["simpson_Dz"] = Dz
    return counts


def unify_party_names(pol: pd.DataFrame) -> pd.DataFrame:
    """Collapse spelling variants of the same party (PARTY_ALIASES) in party_final.

    'Les Engagés'/'LE' and '50PLUS'/'50+' are each one party recorded under two
    labels. Left split, they are counted as two distinct parties, which (e.g.)
    double-counts their seats in reflective_gov_share. Returns a copy.
    """
    pol = pol.copy()
    pol["party_final"] = pol["party_final"].replace(PARTY_ALIASES)
    return pol


def add_lrgen_cat2(pol: pd.DataFrame) -> pd.DataFrame:
    """Add lrgen_cat2, the binary left/right ideology variable (S = 2).

    Starts from lrgen_cat and reassigns the 'center' parties in LRGEN2_RECODE to a
    side (or NaN). Run unify_party_names() first so the recode keys match. Fails
    loudly if any 'center' party is left unmapped, so a new party in future data
    cannot silently drop out of the left/right split. Returns a copy.
    """
    pol = pol.copy()
    keys = list(zip(pol["country"], pol["party_final"]))
    in_map = np.array([k in LRGEN2_RECODE for k in keys])
    mapped = np.array([LRGEN2_RECODE.get(k, np.nan) for k in keys], dtype=object)
    pol["lrgen_cat2"] = np.where(in_map, mapped, pol["lrgen_cat"].to_numpy())
    leftover = (pol.loc[pol["lrgen_cat2"] == "center", ["country", "party_final"]]
                   .drop_duplicates())
    if len(leftover):
        bad = list(map(tuple, leftover.values))
        raise ValueError(f"lrgen_cat2 still has 'center' for {bad}; add them to LRGEN2_RECODE.")
    return pol


def political_dz(pol: pd.DataFrame, group_keys: list,
                 actor_weight: str = "nr_actors") -> pd.DataFrame:
    """Government/opposition Simpson's Dz per grain.

    Government and opposition actors are counted (weighted by actor_weight, since a
    party-in-article row can carry several actors), pooled over group_keys, then
    indexed with S = 2. Dz is the open-diversity score (1 at a 50/50
    government/opposition split). Returns group_keys + n_gov + n_opp + n_actors +
    p_gov + simpson_Dz.
    """
    w = pol[actor_weight].fillna(1) if actor_weight in pol.columns else 1.0
    piv = (pol.assign(_w=w)
              .pivot_table(index=group_keys, columns="government",
                           values="_w", aggfunc="sum", fill_value=0)
              .reindex(columns=[1, 0], fill_value=0))   # [government, opposition]
    Dz, n = diversity_from_counts(piv.to_numpy())
    out = piv.reset_index()
    out.columns = list(out.columns[:-2]) + ["n_gov", "n_opp"]
    out["n_actors"] = n.astype(int)
    out["p_gov"] = np.where(n > 0, out["n_gov"] / n, np.nan)
    out["simpson_Dz"] = Dz
    return out


def ideological_dz(pol: pd.DataFrame, group_keys: list,
                   actor_weight: str = "nr_actors") -> pd.DataFrame:
    """Left/right Simpson's Dz per grain (ideological diversity, S = 2).

    Mirrors political_dz but pools actors by lrgen_cat2 (left vs right) instead of
    government/opposition, so it needs add_lrgen_cat2() to have run. Rows whose
    lrgen_cat2 is NaN (unplaceable parties, e.g. 50+) are excluded. Dz is the open
    score, 1 at a 50/50 left/right split. Returns group_keys + n_left + n_right +
    n_actors + p_right + simpson_Dz.
    """
    d = pol[pol["lrgen_cat2"].isin(LRGEN2_ORDER)]
    w = d[actor_weight].fillna(1) if actor_weight in d.columns else 1.0
    piv = (d.assign(_w=w)
              .pivot_table(index=group_keys, columns="lrgen_cat2",
                           values="_w", aggfunc="sum", fill_value=0)
              .reindex(columns=LRGEN2_ORDER, fill_value=0))   # [left, right]
    Dz, n = diversity_from_counts(piv.to_numpy())
    out = piv.reset_index()
    out.columns = list(out.columns[:-2]) + ["n_left", "n_right"]
    out["n_actors"] = n.astype(int)
    out["p_right"] = np.where(n > 0, out["n_right"] / n, np.nan)
    out["simpson_Dz"] = Dz
    return out


def reflective_gov_share(pol: pd.DataFrame) -> pd.Series:
    """Per-country government share of parliamentary seats (reflective benchmark).

    Built from the distinct parties in the data (one seat share per party), so it
    reflects the pre-election seat distribution among parties that appear in
    coverage. Indexed by country.
    """
    parties = (pol.dropna(subset=["party_final"])
                  .drop_duplicates(["country", "party_final"]))
    num = parties[parties["government"] == 1].groupby("country")["seats"].sum()
    den = parties.groupby("country")["seats"].sum()
    return (num / den).rename("reflective_gov_share")
